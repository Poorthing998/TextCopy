"""
Storage functionality for TextCopy - saves captures to various formats
"""
import logging
from pathlib import Path
from datetime import datetime
from typing import Optional
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from .capture import CapturedText


class StorageManager:
    """Manages storage of captured text to various formats"""

    def __init__(self, config):
        """
        Initialize StorageManager

        Args:
            config: Configuration object
        """
        self.config = config
        self.logger = logging.getLogger(__name__)
        self.capture_count = 0  # Simple counter for numbering captures

    def save_capture(self, capture: CapturedText) -> bool:
        """
        Save capture to configured format(s)

        Args:
            capture: CapturedText object to save

        Returns:
            True if successful, False otherwise
        """
        output_format = self.config.get('output_format', 'both')
        word_success = True
        text_success = True

        try:
            if output_format in ['word', 'both']:
                word_success = self.save_to_word(capture)
                if not word_success:
                    self.logger.error("Failed to save to Word document")

            if output_format in ['txt', 'both']:
                text_success = self.save_to_text(capture)
                if not text_success:
                    self.logger.error("Failed to save to text file")

            return word_success and text_success

        except Exception as e:
            self.logger.error(f"Error saving capture: {e}", exc_info=True)
            return False

    def save_to_word(self, capture: CapturedText) -> bool:
        """
        Save capture to Word document - SIMPLIFIED

        Args:
            capture: CapturedText object to save

        Returns:
            True if successful, False otherwise
        """
        word_path = self.config.get_word_output_path()

        # Try to save, with fallback if file is locked (Windows)
        use_backup = self.config.get('use_backup_on_lock', True)
        max_attempts = 3 if use_backup else 1

        for attempt in range(max_attempts):
            try:
                # Determine which file to use
                if attempt == 0:
                    current_path = word_path
                else:
                    backup_name = word_path.stem + f"_backup{attempt}" + word_path.suffix
                    current_path = word_path.parent / backup_name

                # Load existing document or create new one
                if current_path.exists():
                    doc = Document(str(current_path))
                    # Count existing numbered items to continue numbering
                    self.capture_count = self._count_captures(doc)
                else:
                    doc = Document()
                    self.capture_count = 0

                # Increment counter for this capture
                self.capture_count += 1

                # SIMPLE FORMAT: Just number and text
                para = doc.add_paragraph()

                # Add number
                number_run = para.add_run(f"{self.capture_count}. ")
                number_run.font.bold = True
                number_run.font.size = Pt(11)

                # Add text
                text_run = para.add_run(capture.text)
                text_run.font.name = "Calibri"
                text_run.font.size = Pt(11)

                # Add blank line
                doc.add_paragraph()

                # Save document
                doc.save(str(current_path))

                if attempt > 0:
                    self.logger.info(f"Saved to backup: {current_path.name}")
                else:
                    self.logger.info(f"✓ Saved capture #{self.capture_count}")

                return True

            except PermissionError:
                if attempt < max_attempts - 1:
                    self.logger.warning(f"File locked, trying backup...")
                    continue
                else:
                    self.logger.error(f"Cannot save - close {word_path.name} in Word")
                    return False

            except Exception as e:
                self.logger.error(f"Error: {e}")
                return False

        return False

    def _count_captures(self, doc: Document) -> int:
        """Count how many captures are already in the document"""
        count = 0
        for para in doc.paragraphs:
            text = para.text.strip()
            if text and text[0].isdigit() and '. ' in text:
                try:
                    num = int(text.split('.')[0])
                    if num > count:
                        count = num
                except:
                    pass
        return count

    def save_to_text(self, capture: CapturedText) -> bool:
        """
        Save capture to text file

        Args:
            capture: CapturedText object to save

        Returns:
            True if successful, False otherwise
        """
        try:
            text_path = self.config.get_text_output_path()

            # Build text entry
            lines = []

            if self.config.get('include_separator', True):
                lines.append('\n' + '=' * 60)

            if self.config.get('include_timestamp', True):
                lines.append(f"[{capture.timestamp.strftime('%Y-%m-%d %H:%M:%S')}]")

            if self.config.get('include_source', True) and capture.source:
                lines.append(f"Source: {capture.source}")

            if lines:
                lines.append('')  # Empty line before text

            lines.append(capture.text)
            lines.append('')  # Empty line after text

            # Append to file
            with open(text_path, 'a', encoding='utf-8') as f:
                f.write('\n'.join(lines))

            self.logger.info(f"Saved to text file: {text_path}")

            return True

        except Exception as e:
            self.logger.error(f"Error saving to text: {e}")
            return False

    def _add_document_header(self, doc: Document) -> None:
        """
        Add header to new Word document

        Args:
            doc: Document object
        """
        # Title
        title = doc.add_heading('TextCopy - Captured Text', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Description
        desc = doc.add_paragraph()
        desc_run = desc.add_run(
            f"Text captures collected with TextCopy\n"
            f"Started: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        desc_run.font.size = Pt(10)
        desc_run.font.italic = True
        desc_run.font.color.rgb = RGBColor(102, 102, 102)
        desc.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Divider
        divider = doc.add_paragraph()
        run = divider.add_run('═' * 60)
        run.font.color.rgb = RGBColor(128, 128, 128)
        divider.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    def get_word_path(self) -> Path:
        """Get path to Word document"""
        return self.config.get_word_output_path()

    def get_text_path(self) -> Path:
        """Get path to text file"""
        return self.config.get_text_output_path()

    def open_captures(self) -> bool:
        """
        Open captures file with default application

        Returns:
            True if successful, False otherwise
        """
        try:
            import platform
            import subprocess

            output_format = self.config.get('output_format', 'both')

            # Determine which file to open
            if output_format in ['word', 'both']:
                file_path = self.get_word_path()
            else:
                file_path = self.get_text_path()

            if not file_path.exists():
                self.logger.warning(f"File does not exist: {file_path}")
                return False

            # Open file with default application
            system = platform.system()

            if system == "Windows":
                subprocess.Popen(['start', str(file_path)], shell=True)
            elif system == "Darwin":  # macOS
                subprocess.Popen(['open', str(file_path)])
            elif system == "Linux":
                subprocess.Popen(['xdg-open', str(file_path)])

            self.logger.info(f"Opened captures file: {file_path}")
            return True

        except Exception as e:
            self.logger.error(f"Error opening captures: {e}")
            return False

    def get_file_stats(self) -> dict:
        """
        Get statistics about capture files

        Returns:
            Dictionary with file statistics
        """
        stats = {}

        try:
            word_path = self.get_word_path()
            if word_path.exists():
                stats['word_size'] = word_path.stat().st_size
                stats['word_modified'] = datetime.fromtimestamp(word_path.stat().st_mtime)

            text_path = self.get_text_path()
            if text_path.exists():
                stats['text_size'] = text_path.stat().st_size
                stats['text_modified'] = datetime.fromtimestamp(text_path.stat().st_mtime)

                # Count lines in text file
                with open(text_path, 'r', encoding='utf-8') as f:
                    stats['text_lines'] = sum(1 for _ in f)

        except Exception as e:
            self.logger.error(f"Error getting file stats: {e}")

        return stats
